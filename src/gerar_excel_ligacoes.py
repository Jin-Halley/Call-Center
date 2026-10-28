# Junta relatórios de ligações do PABX (.csv/.txt), remove duplicadas e gera um Excel formatado
# Uso: python gerar_excel_ligacoes.py arquivo1.txt arquivo2.csv saida.xlsx   (precisa: pip install openpyxl)

import csv, io, sys, datetime as dt, collections
def ler(path):
    raw=open(path,'rb').read()
    for enc in ('utf-8-sig','cp1252','latin-1'):
        try: txt=raw.decode(enc); break
        except: pass
    out=[]
    for line in txt.splitlines():
        s=line.strip()
        if not s: continue
        if s.startswith('"') and s.endswith('"') and s.count('""')>0:  # linha inteira entre aspas
            s=s[1:-1].replace('""','"')
        row=next(csv.reader([s]))
        if row and row[0].strip().lower()=='data': continue
        out.append(tuple(c.strip() for c in row[:14]))
    return out

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as L

rows=[]
seen=set()
for p in sys.argv[1:-1]:
    for r in ler(p):
        if r in seen: continue
        seen.add(r); rows.append(r)
rows.sort(key=lambda r:(dt.datetime.strptime(r[0],'%d/%m/%Y'),r[1]))

def tel(n):
    d=''.join(ch for ch in n if ch.isdigit())
    if d.startswith('55') and len(d)>=12: d=d[2:]
    if len(d)==11: return f'({d[:2]}) {d[2:7]}-{d[7:]}'
    if len(d)==10: return f'({d[:2]}) {d[2:6]}-{d[6:]}'
    return n
def hms(s):
    if not s: return None
    h,m,se=map(int,s.split(':')); return dt.time(h,m,se)
SENT={'A originar':'Saída','A finalizar':'Entrada','Intra BG':'Intra BG'}

F="Arial"; fill=lambda c: PatternFill("solid",start_color=c,end_color=c)
NAVY="1F3864"; BLUE="2E75B6"; GREEN="1E7B3A"; RED="B42318"; G="6B7785"; Z="F5F8FC"
t=Side(style="thin",color="D5DCE6"); B=Border(left=t,right=t,top=t,bottom=t)
wb=Workbook()

# ---------- Chamadas
C=wb.active; C.title="Chamadas"; C.sheet_view.showGridLines=False
H=["Data","Hora","Sentido","Tipo original","Ramal","Número de origem","Número chamado","Ramal destino","Atendida","Duração","Espera","Duração (s)"]
for k,h in enumerate(H,1):
    c=C.cell(1,k,h); c.font=Font(name=F,bold=True,color="FFFFFF"); c.fill=fill(NAVY); c.alignment=Alignment(horizontal="center",vertical="center")
for i,r in enumerate(rows,2):
    d=dt.datetime.strptime(r[0],'%d/%m/%Y'); h=dt.datetime.strptime(r[1],'%H:%M:%S').time()
    ramal=int(r[4]) if r[4] else (int(r[7]) if r[7] else None)
    vals=[d,h,SENT.get(r[2],r[2]),r[2],ramal,tel(r[3]),tel(r[6]),int(r[7]) if r[7] else None,r[9],hms(r[10]),hms(r[11])]
    for k,v in enumerate(vals,1):
        c=C.cell(i,k,v); c.font=Font(name=F,size=10); c.border=B
        c.alignment=Alignment(horizontal="left" if k in(6,7) else "center")
    C.cell(i,1).number_format='dd/mm/yyyy'; C.cell(i,2).number_format='hh:mm:ss'
    for k in (10,11): C.cell(i,k).number_format='[h]:mm:ss'
    c=C.cell(i,12,f'=IF(J{i}="","",ROUND(J{i}*86400,0))'); c.font=Font(name=F,size=10,color=G); c.border=B; c.alignment=Alignment(horizontal="center")
n=len(rows); last=n+1
for k,w in enumerate([12,10,11,13,8,19,19,13,10,10,9,11],1): C.column_dimensions[L(k)].width=w
tb=Table(displayName="Chamadas",ref=f"A1:L{last}"); tb.tableStyleInfo=TableStyleInfo(name="TableStyleLight9",showRowStripes=True); C.add_table(tb)
C.freeze_panes="A2"
C.conditional_formatting.add(f"I2:I{last}",FormulaRule(formula=['$I2="Sim"'],font=Font(name=F,bold=True,color=GREEN),fill=fill("DFF3E4")))
C.conditional_formatting.add(f"I2:I{last}",FormulaRule(formula=['$I2="Não"'],font=Font(name=F,bold=True,color=RED),fill=fill("FBE3E1")))
C.conditional_formatting.add(f"J2:J{last}",FormulaRule(formula=['$J2>=TIME(0,5,0)'],font=Font(name=F,bold=True,color="8A5A00"),fill=fill("FFF1D6")))

# ---------- Resumo
R=wb.create_sheet("Resumo",0); R.sheet_view.showGridLines=False
for k,w in enumerate([2,14,13,13,13,13,15,15,3,10,13,13],1): R.column_dimensions[L(k)].width=w
datas=sorted({r[0] for r in rows})
R.merge_cells("B1:L1"); c=R["B1"]; c.value=f"RELATÓRIO DE LIGAÇÕES  •  {datas[0]}" + (f" a {datas[-1]}" if len(datas)>1 else "")
c.font=Font(name=F,size=16,bold=True,color="FFFFFF"); c.fill=fill(NAVY); c.alignment=Alignment(vertical="center",indent=1); R.row_dimensions[1].height=34
R.merge_cells("B2:L2"); R["B2"].value="Fonte: registro de chamadas do PABX. Tudo nesta aba é calculado a partir da aba Chamadas."; R["B2"].font=Font(name=F,size=9,italic=True,color=G)
K=[("Ligações",'=ROWS(Chamadas[Data])','0'),("Atendidas",'=COUNTIF(Chamadas[Atendida],"Sim")','0'),("Não atendidas",'=COUNTIF(Chamadas[Atendida],"Não")','0'),
   ("% atendidas",'=IFERROR(C5/B5,0)','0.0%'),("Tempo falado",'=SUM(Chamadas[Duração])','[h]:mm:ss'),("Média atendida",'=IFERROR(AVERAGEIF(Chamadas[Atendida],"Sim",Chamadas[Duração]),0)','[h]:mm:ss'),("Maior ligação",'=MAX(Chamadas[Duração])','[h]:mm:ss')]
for j,(lab,f,fmt) in enumerate(K):
    col=2+j
    a=R.cell(4,col,lab); a.font=Font(name=F,size=9,bold=True,color=G); a.alignment=Alignment(horizontal="center"); a.fill=fill(Z); a.border=B
    v=R.cell(5,col,f); v.font=Font(name=F,size=15,bold=True,color=NAVY); v.alignment=Alignment(horizontal="center",vertical="center"); v.number_format=fmt; v.fill=fill(Z); v.border=B
R.row_dimensions[5].height=30
for col,colr in ((3,GREEN),(4,RED)): R.cell(5,col).font=Font(name=F,size=15,bold=True,color=colr)

def head(row,cols,start=2):
    for k,h in enumerate(cols):
        c=R.cell(row,start+k,h); c.font=Font(name=F,bold=True,color="FFFFFF",size=10); c.fill=fill(BLUE); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=B
    R.row_dimensions[row].height=30
def sec(row,txt,a="B",b="H"):
    R.merge_cells(f"{a}{row}:{b}{row}"); c=R[f"{a}{row}"]; c.value=txt; c.font=Font(name=F,bold=True,size=11,color=NAVY)

# por ramal
sec(7,"POR RAMAL")
head(8,["Ramal","Ligações","Atendidas","Não atendidas","% atendidas","Tempo falado","Tempo médio (atendidas)"])
ramais=sorted({ (int(r[4]) if r[4] else int(r[7])) for r in rows if r[4] or r[7]})
r0=9
for i,ra in enumerate(ramais):
    r=r0+i
    vals=[ra,f'=COUNTIF(Chamadas[Ramal],B{r})',f'=COUNTIFS(Chamadas[Ramal],B{r},Chamadas[Atendida],"Sim")',f'=COUNTIFS(Chamadas[Ramal],B{r},Chamadas[Atendida],"Não")',
          f'=IFERROR(D{r}/C{r},0)',f'=SUMIF(Chamadas[Ramal],B{r},Chamadas[Duração])',f'=IFERROR(AVERAGEIFS(Chamadas[Duração],Chamadas[Ramal],B{r},Chamadas[Atendida],"Sim"),0)']
    for k,v in enumerate(vals):
        c=R.cell(r,2+k,v); c.font=Font(name=F,size=10,bold=(k==0)); c.border=B; c.alignment=Alignment(horizontal="center"); c.fill=fill("FFFFFF" if i%2==0 else Z)
    R.cell(r,6).number_format='0%'; R.cell(r,7).number_format='[h]:mm:ss'; R.cell(r,8).number_format='[h]:mm:ss'
rt=r0+len(ramais)
R.cell(rt,2,"TOTAL")
for col,f,fmt in ((3,f'=SUM(C{r0}:C{rt-1})','0'),(4,f'=SUM(D{r0}:D{rt-1})','0'),(5,f'=SUM(E{r0}:E{rt-1})','0'),(6,f'=IFERROR(D{rt}/C{rt},0)','0%'),(7,f'=SUM(G{r0}:G{rt-1})','[h]:mm:ss'),(8,'=G5','[h]:mm:ss')):
    R.cell(rt,col,f).number_format=fmt
for col in range(2,9):
    c=R.cell(rt,col); c.font=Font(name=F,bold=True,color="FFFFFF"); c.fill=fill(NAVY); c.alignment=Alignment(horizontal="center"); c.border=B
R.conditional_formatting.add(f"C{r0}:C{rt-1}",DataBarRule(start_type='num',start_value=0,end_type='max',color="8FB4DC"))
R.conditional_formatting.add(f"G{r0}:G{rt-1}",DataBarRule(start_type='num',start_value=0,end_type='max',color="9FD3AE"))

# por sentido
sec(7,"POR TIPO","J","L")
for k,h in enumerate(["Tipo","Ligações","Atendidas"]):
    c=R.cell(8,10+k,h); c.font=Font(name=F,bold=True,color="FFFFFF",size=10); c.fill=fill(BLUE); c.alignment=Alignment(horizontal="center",vertical="center"); c.border=B
for i,s in enumerate(["Saída","Entrada","Intra BG"]):
    r=9+i
    for k,v in enumerate([s,f'=COUNTIF(Chamadas[Sentido],J{r})',f'=COUNTIFS(Chamadas[Sentido],J{r},Chamadas[Atendida],"Sim")']):
        c=R.cell(r,10+k,v); c.font=Font(name=F,size=10,bold=(k==0)); c.border=B; c.alignment=Alignment(horizontal="center")


# por hora
hs=rt+3
sec(hs-1,"POR HORA DO DIA")
head(hs,["Hora","Ligações","Atendidas","Não atendidas","Tempo falado"])
horas=list(range(min(int(r[1][:2]) for r in rows),max(int(r[1][:2]) for r in rows)+1))
for i,hh in enumerate(horas):
    r=hs+1+i
    vals=[f"{hh:02d}h",f'=COUNTIFS(Chamadas[Hora],">="&TIME({hh},0,0),Chamadas[Hora],"<"&TIME({hh+1},0,0))',
          f'=COUNTIFS(Chamadas[Hora],">="&TIME({hh},0,0),Chamadas[Hora],"<"&TIME({hh+1},0,0),Chamadas[Atendida],"Sim")',f'=C{r}-D{r}',
          f'=SUMIFS(Chamadas[Duração],Chamadas[Hora],">="&TIME({hh},0,0),Chamadas[Hora],"<"&TIME({hh+1},0,0))']
    for k,v in enumerate(vals):
        c=R.cell(r,2+k,v); c.font=Font(name=F,size=10,bold=(k==0)); c.border=B; c.alignment=Alignment(horizontal="center"); c.fill=fill("FFFFFF" if i%2==0 else Z)
    R.cell(r,6).number_format='[h]:mm:ss'
he=hs+len(horas)

ch=BarChart(); ch.type="col"; ch.grouping="stacked"; ch.overlap=100; ch.title="Ligações por hora"; ch.height=7.5; ch.width=17
ch.add_data(Reference(R,min_col=4,max_col=5,min_row=hs,max_row=he),titles_from_data=True)
ch.set_categories(Reference(R,min_col=2,min_row=hs+1,max_row=he))
ch.series[0].graphicalProperties.solidFill=GREEN; ch.series[1].graphicalProperties.solidFill="E06666"
ch.legend.position="b"; ch.y_axis.delete=False; ch.x_axis.delete=False; ch.y_axis.majorGridlines=None
R.add_chart(ch,f"H{hs}")
ch2=BarChart(); ch2.type="bar"; ch2.title="Ligações por ramal"; ch2.height=7.5; ch2.width=17; ch2.grouping="stacked"; ch2.overlap=100
ch2.add_data(Reference(R,min_col=4,max_col=5,min_row=8,max_row=rt-1),titles_from_data=True)
ch2.set_categories(Reference(R,min_col=2,min_row=r0,max_row=rt-1))
ch2.series[0].graphicalProperties.solidFill=GREEN; ch2.series[1].graphicalProperties.solidFill="E06666"
ch2.legend.position="b"; ch2.y_axis.delete=False; ch2.x_axis.delete=False; ch2.y_axis.majorGridlines=None
R.add_chart(ch2,f"H{he+3}")

# números mais chamados
cnt=collections.Counter(tel(r[6]) for r in rows if r[2]!='A finalizar')
top=[x for x in cnt.most_common() if x[1]>=2][:15]
ts=he+3
sec(ts-1,"NÚMEROS MAIS CHAMADOS (2 vezes ou mais)","B","F")
head(ts,["Número","Tentativas","Atendidas","Tempo falado"])
for i,(num,_) in enumerate(top):
    r=ts+1+i
    vals=[num,f'=COUNTIF(Chamadas[Número chamado],B{r})',f'=COUNTIFS(Chamadas[Número chamado],B{r},Chamadas[Atendida],"Sim")',f'=SUMIF(Chamadas[Número chamado],B{r},Chamadas[Duração])']
    for k,v in enumerate(vals):
        c=R.cell(r,2+k,v); c.font=Font(name=F,size=10,bold=(k==0)); c.border=B; c.alignment=Alignment(horizontal="center"); c.fill=fill("FFFFFF" if i%2==0 else Z)
    R.cell(r,5).number_format='[h]:mm:ss'
R.column_dimensions['B'].width=17
R.page_setup.orientation="portrait"; R.page_setup.fitToWidth=1; R.page_setup.fitToHeight=0; R.sheet_properties.pageSetUpPr.fitToPage=True
C.page_setup.orientation="landscape"; C.page_setup.fitToWidth=1; C.page_setup.fitToHeight=0; C.sheet_properties.pageSetUpPr.fitToPage=True; C.print_title_rows="1:1"
out=sys.argv[-1]; wb.save(out); print(out,n,'ligações',len(ramais),'ramais')
